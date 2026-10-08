import logging
import math
from typing import Any
import fitz
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from app.services.pdf_to_pptx.coordinates import SlideGeometry
from app.services.pdf_to_pptx.fonts import (
    clean_xml_string,
    is_bold,
    is_garbage_or_symbol_glyph,
    is_italic,
    normalize_font_name,
    parse_color,
)

logger = logging.getLogger(__name__)


def _detect_alignment(lines: list[dict[str, Any]]) -> PP_ALIGN:
    """Infers text block alignment based on line bounding boxes."""
    if len(lines) <= 1:
        return PP_ALIGN.LEFT

    x0_coords = [ln["bbox"][0] for ln in lines if ln.get("bbox")]
    x1_coords = [ln["bbox"][2] for ln in lines if ln.get("bbox")]
    centers = [(ln["bbox"][0] + ln["bbox"][2]) / 2.0 for ln in lines if ln.get("bbox")]

    if not x0_coords or not x1_coords:
        return PP_ALIGN.LEFT

    var_left = max(x0_coords) - min(x0_coords)
    var_center = max(centers) - min(centers)
    var_right = max(x1_coords) - min(x1_coords)

    # If center points are aligned tightly
    if var_center < 5.0 and var_center < var_left:
        return PP_ALIGN.CENTER
    # If right margins are aligned tightly
    if var_right < 5.0 and var_right < var_left:
        return PP_ALIGN.RIGHT

    return PP_ALIGN.LEFT


def extract_and_add_text(
    page: fitz.Page,
    slide: Any,
    geom: SlideGeometry,
    excluded_rects: list[fitz.Rect] | None = None,
    default_font: str = "Calibri",
) -> int:
    """Extracts digital text and places it as native, editable PowerPoint text boxes.

    Preserves exact coordinates, typography, calculated line spacing in points,
    natural paragraph grouping, and prevents unintended word wrapping.
    """
    excluded = excluded_rects or []
    try:
        page_dict = page.get_text("dict", flags=fitz.TEXTFLAGS_SEARCH)
    except Exception as exc:
        logger.warning("Could not read text dict on page %d: %s", page.number + 1, exc)
        return 0

    blocks = page_dict.get("blocks", [])
    if not blocks:
        return 0

    added_count = 0

    for block in blocks:
        if block.get("type") != 0:  # 0 is text block
            continue

        bbox_raw = block.get("bbox")
        if not bbox_raw or len(bbox_raw) != 4:
            continue

        b_rect = fitz.Rect(bbox_raw)
        if b_rect.width <= 1.0 or b_rect.height <= 1.0:
            continue

        # Skip text blocks entirely inside excluded tables
        if any(b_rect in ex_rect or ex_rect.contains(b_rect) for ex_rect in excluded):
            continue

        lines = block.get("lines", [])
        if not lines:
            continue

        # Pre-filter lines and spans: remove non-text symbol glyphs and clean XML
        valid_lines: list[dict[str, Any]] = []
        for line in lines:
            valid_spans: list[dict[str, Any]] = []
            for span in line.get("spans", []):
                raw_text = span.get("text", "")
                font_raw = span.get("font", "")

                # Reject non-text graphical icons or unmapped PUA codes
                if is_garbage_or_symbol_glyph(raw_text, font_raw):
                    continue

                cleaned_text = clean_xml_string(raw_text)
                if not cleaned_text.strip():
                    continue

                span_copy = dict(span)
                span_copy["text"] = cleaned_text
                valid_spans.append(span_copy)

            if valid_spans:
                line_copy = dict(line)
                line_copy["spans"] = valid_spans
                valid_lines.append(line_copy)

        if not valid_lines:
            continue

        left, top, width, height = geom.to_pptx_coords(b_rect.x0, b_rect.y0, b_rect.x1, b_rect.y1)

        # For single line headings/labels, do not wrap text so words never prematurely wrap
        is_single_line = (len(valid_lines) <= 1)
        if is_single_line:
            # Add safety margin so font metric variances don't clip text
            width = Emu(int(width * 1.25))
        else:
            width = Emu(int(width * 1.15))

        try:
            tx_box = slide.shapes.add_textbox(left, top, width, height)
            tf = tx_box.text_frame
            # Single line text boxes should not wrap; multi-line paragraphs should wrap
            tf.word_wrap = not is_single_line
            # Zero margins ensure text position matches exact PDF coordinates
            tf.margin_left = Pt(0)
            tf.margin_right = Pt(0)
            tf.margin_top = Pt(0)
            tf.margin_bottom = Pt(0)

            # Rotation from the first line direction vector
            first_dir = valid_lines[0].get("dir", (1.0, 0.0))
            if len(first_dir) == 2 and (abs(first_dir[1]) > 0.05 or first_dir[0] < 0.95):
                angle_deg = math.degrees(math.atan2(first_dir[1], first_dir[0]))
                if abs(angle_deg) > 1.0:
                    tx_box.rotation = angle_deg

            align = _detect_alignment(valid_lines)

            # Calculate actual line spacing from vertical delta
            line_spacing_pt: float | None = None
            if len(valid_lines) >= 2:
                delta_y = valid_lines[1]["bbox"][1] - valid_lines[0]["bbox"][1]
                if delta_y > 4.0:
                    line_spacing_pt = round(delta_y, 1)

            current_p = tf.paragraphs[0]
            current_p.alignment = align
            if line_spacing_pt is not None:
                current_p.line_spacing = Pt(line_spacing_pt)

            prev_line_bottom: float | None = None
            prev_font_size: float = 12.0

            for line_idx, line in enumerate(valid_lines):
                line_bbox = line.get("bbox", (0, 0, 0, 0))
                line_top = line_bbox[1]
                spans = line.get("spans", [])

                # Determine if a new paragraph should be started or continued
                if line_idx > 0:
                    y_gap = (line_top - prev_line_bottom) if prev_line_bottom is not None else 0
                    is_new_para = y_gap > (prev_font_size * 1.6)

                    if is_new_para:
                        current_p = tf.add_paragraph()
                        current_p.alignment = align
                        current_p.space_before = Pt(3)
                        if line_spacing_pt is not None:
                            current_p.line_spacing = Pt(line_spacing_pt)
                    else:
                        br_run = current_p.add_run()
                        br_run.text = "\n"

                for span_idx, span in enumerate(spans):
                    text_content = span.get("text", "")
                    if not text_content:
                        continue

                    run = current_p.add_run()
                    run.text = text_content

                    font_raw = span.get("font", "")
                    run.font.name = normalize_font_name(font_raw, default_font)

                    size_val = float(span.get("size", 12.0) or 12.0)
                    size_val = max(4.0, min(144.0, size_val))
                    run.font.size = Pt(size_val)
                    prev_font_size = size_val

                    flags = int(span.get("flags", 0) or 0)
                    run.font.bold = is_bold(flags, font_raw)
                    run.font.italic = is_italic(flags, font_raw)
                    run.font.color.rgb = parse_color(span.get("color", 0))

                prev_line_bottom = line_bbox[3]

            added_count += 1
        except Exception as exc:
            logger.debug("Failed adding text box on slide %d: %s", page.number + 1, exc)

    logger.debug("Added %d editable text boxes to slide %d", added_count, page.number + 1)
    return added_count
