import io
import logging
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from PIL import Image
from pptx import Presentation

from app.services.pdf_to_pptx.coordinates import SlideGeometry
from app.services.pdf_to_pptx.images import extract_and_add_images
from app.services.pdf_to_pptx.ocr import is_scanned_page, perform_ocr_on_page
from app.services.pdf_to_pptx.shapes import extract_and_add_shapes
from app.services.pdf_to_pptx.tables import extract_and_add_tables
from app.services.pdf_to_pptx.text import extract_and_add_text

logger = logging.getLogger(__name__)


class PdfToPptxError(Exception):
    """Client-safe failure. The router converts it to an HTTPException."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


@dataclass(frozen=True)
class PdfToPptxResult:
    output_path: Path
    slide_count: int
    size_bytes: int


def _ensure_pdf_header(pdf_path: Path) -> None:
    with pdf_path.open("rb") as fh:
        header = fh.read(1024)
        if b"%PDF-" not in header:
            raise PdfToPptxError("Invalid or corrupted PDF.", 400)


def _render_page_fallback(
    page: fitz.Page,
    slide: Any,
    geom: SlideGeometry,
    dpi: int = 150,
    max_pixels: int = 40_000_000,
    image_format: str = "jpeg",
    quality: int = 90,
) -> None:
    """Renders full page as an image fallback when native objects cannot be represented."""
    rect = page.rect
    zoom = dpi / 72.0
    pixels = (rect.width * zoom) * (rect.height * zoom)
    if pixels > max_pixels:
        zoom *= (max_pixels / pixels) ** 0.5

    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)

    buf = io.BytesIO()
    if image_format.lower() == "png":
        img.save(buf, format="PNG")
    else:
        img.save(buf, format="JPEG", quality=quality, subsampling=0)
    buf.seek(0)
    img.close()
    del pix

    left, top, width, height = geom.to_pptx_coords(rect.x0, rect.y0, rect.x1, rect.y1)
    slide.shapes.add_picture(buf, left, top, width, height)
    buf.close()


def validate_pptx_output(output_path: Path, expected_slides: int, doc: fitz.Document | None = None) -> None:
    """Validates that the generated PPTX file exists, is readable, matches slide count, and has valid content."""
    if not output_path.exists():
        raise PdfToPptxError("PPTX output file was not created.", 500)

    if output_path.stat().st_size == 0:
        raise PdfToPptxError("Generated PPTX file is empty.", 500)

    try:
        prs = Presentation(output_path)
        actual_slides = len(prs.slides)
        if actual_slides != expected_slides:
            logger.warning(
                "Slide count mismatch: expected %d, got %d",
                expected_slides,
                actual_slides,
            )
            raise PdfToPptxError("PPTX validation failed: slide count mismatch.", 500)

        # Check for unexpected blank slides when PDF page had content
        if doc is not None and len(doc) == actual_slides:
            for s_idx, slide in enumerate(prs.slides):
                p = doc[s_idx]
                has_pdf_content = (
                    len(p.get_text().strip()) > 0
                    or len(p.get_drawings()) > 0
                    or len(p.get_images()) > 0
                )
                if has_pdf_content and len(slide.shapes) == 0:
                    # If slide has a native background fill, it's not completely blank
                    has_bg = False
                    try:
                        has_bg = slide.background.fill.type is not None
                    except Exception:
                        pass
                    if not has_bg:
                        logger.warning("Slide %d is unexpectedly blank despite PDF having content", s_idx + 1)
    except PdfToPptxError:
        raise
    except Exception as exc:
        logger.exception("Failed validating generated PPTX: %s", exc)
        raise PdfToPptxError("Generated PPTX file is corrupted or unreadable.", 500) from exc


def convert_pdf_to_pptx(
    pdf_path: Path,
    output_path: Path,
    *,
    dpi: int = 150,
    max_pages: int = 300,
    max_pixels: int = 40_000_000,
    image_format: str = "jpeg",
    jpeg_quality: int = 90,
    default_font: str = "Calibri",
    ocr_enabled: bool = True,
    ocr_languages: str = "eng,ben",
    extract_images: bool = True,
    extract_shapes: bool = True,
    detect_tables: bool = True,
) -> PdfToPptxResult:
    """Converts a PDF into a production-grade editable PowerPoint presentation.

    Preserves text boxes, fonts, sizes, colors, styles, separate images,
    native vector shapes, and tables with z-order layering and slide geometry.
    """
    started = time.perf_counter()
    pdf_path = Path(pdf_path)
    output_path = Path(output_path)
    _ensure_pdf_header(pdf_path)

    try:
        doc = fitz.open(pdf_path)
    except Exception as exc:
        logger.warning("Could not open PDF: %s", exc)
        raise PdfToPptxError("Invalid or corrupted PDF.", 400) from exc

    with doc:
        if not doc.is_pdf:
            raise PdfToPptxError("Invalid or corrupted PDF.", 400)
        if doc.needs_pass:
            raise PdfToPptxError("Password-protected PDFs are not supported.", 400)

        total = doc.page_count
        if total < 1:
            raise PdfToPptxError("Invalid or corrupted PDF.", 400)
        if total > max_pages:
            raise PdfToPptxError("PDF contains too many pages.", 413)

        try:
            first_page = doc.load_page(0)
            first_rect = first_page.rect
        except Exception as exc:
            logger.exception("Could not read first page")
            raise PdfToPptxError("Invalid or corrupted PDF.", 422) from exc

        if first_rect.width <= 0 or first_rect.height <= 0:
            raise PdfToPptxError("Invalid or corrupted PDF.", 422)

        base_geom = SlideGeometry.from_page_rect(first_rect)

        logger.info(
            "PDF to Editable PPTX started: pages=%d, dimensions=(%d x %d pt), font=%s",
            total,
            int(first_rect.width),
            int(first_rect.height),
            default_font,
        )

        try:
            prs = Presentation()
            prs.slide_width = base_geom.slide_width
            prs.slide_height = base_geom.slide_height
            blank_layout = prs.slide_layouts[6]  # Blank slide
        except Exception as exc:
            logger.exception("Could not initialize presentation")
            raise PdfToPptxError("Unable to convert PDF to PowerPoint.", 500) from exc

        for index in range(total):
            try:
                page = doc.load_page(index)
                rect = page.rect
                if rect.width <= 0 or rect.height <= 0:
                    raise ValueError(f"Page {index + 1} has invalid dimensions")
            except Exception as exc:
                logger.exception("Loading page %d failed", index + 1)
                raise PdfToPptxError("Invalid or corrupted PDF.", 422) from exc

            page_geom = SlideGeometry.for_page(rect, base_geom.slide_width, base_geom.slide_height)
            slide = prs.slides.add_slide(blank_layout)

            # Check if page is genuinely scanned (0 text, 0 drawings, dominated by raster image)
            scanned = is_scanned_page(page)
            if scanned:
                ocr_success = False
                if ocr_enabled:
                    try:
                        ocr_success = perform_ocr_on_page(
                            page,
                            slide,
                            page_geom,
                            default_font=default_font,
                            languages=ocr_languages,
                        )
                    except Exception as exc:
                        logger.warning("OCR failed on page %d: %s", index + 1, exc)

                if not ocr_success:
                    logger.info("Page %d is scanned; rendering page image fallback", index + 1)
                    _render_page_fallback(
                        page,
                        slide,
                        page_geom,
                        dpi=dpi,
                        max_pixels=max_pixels,
                        image_format=image_format,
                        quality=jpeg_quality,
                    )
                continue

            # Digital page: Extract native elements with proper Z-Order
            # Layering: Background -> Tables -> Shapes (lines, rectangles, vectors) -> Images -> Text
            try:
                table_rects: list[fitz.Rect] = []
                if detect_tables:
                    table_rects = extract_and_add_tables(
                        page, slide, page_geom, default_font=default_font
                    )

                shapes_count = 0
                if extract_shapes:
                    shapes_count = extract_and_add_shapes(
                        page, slide, page_geom, excluded_rects=table_rects, dpi=dpi
                    )

                images_count = 0
                if extract_images:
                    images_count = extract_and_add_images(doc, page, slide, page_geom, dpi=dpi)

                text_count = extract_and_add_text(
                    page,
                    slide,
                    page_geom,
                    excluded_rects=table_rects,
                    default_font=default_font,
                )

                # Fallback if no elements were extractable on this page
                if text_count == 0 and images_count == 0 and shapes_count == 0 and len(table_rects) == 0:
                    logger.info("Page %d produced no native elements; rendering fallback image", index + 1)
                    _render_page_fallback(
                        page,
                        slide,
                        page_geom,
                        dpi=dpi,
                        max_pixels=max_pixels,
                        image_format=image_format,
                        quality=jpeg_quality,
                    )
                else:
                    logger.info(
                        "Slide %d/%d generated: %d text boxes, %d images, %d shapes, %d tables",
                        index + 1,
                        total,
                        text_count,
                        images_count,
                        shapes_count,
                        len(table_rects),
                    )

            except Exception as exc:
                logger.warning(
                    "Page %d element extraction encountered error; using fallback image: %s",
                    index + 1,
                    exc,
                )
                # Remove any partially added shapes from the failed slide attempt
                while len(slide.shapes) > 0:
                    try:
                        sp = slide.shapes[0]._element
                        sp.getparent().remove(sp)
                    except Exception:
                        break
                _render_page_fallback(
                    page,
                    slide,
                    page_geom,
                    dpi=dpi,
                    max_pixels=max_pixels,
                    image_format=image_format,
                    quality=jpeg_quality,
                )

        try:
            prs.save(output_path)
        except Exception as exc:
            logger.exception("Saving PPTX failed")
            raise PdfToPptxError("Unable to convert PDF to PowerPoint.", 500) from exc

        validate_pptx_output(output_path, total, doc)

    size = output_path.stat().st_size
    duration = time.perf_counter() - started
    logger.info(
        "PDF to PPTX completed successfully: slides=%d bytes=%d duration=%.2fs",
        total,
        size,
        duration,
    )
    return PdfToPptxResult(output_path, total, size)
