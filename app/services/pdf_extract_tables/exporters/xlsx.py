from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.pdf_extract_tables.exporters.base import MAX_CELL_CHARS, ExportedFile
from app.services.pdf_extract_tables.models import Table, TableDocument
from app.services.pdf_extract_tables.normalizer import _is_numeric

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_SHEET_TITLE = 31
MIN_COL_WIDTH = 10
MAX_COL_WIDTH = 60

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\\x00-\x1f]")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TITLE_FILL = PatternFill("solid", fgColor="E2E8F0")
_TITLE_FONT = Font(name="Calibri", size=12, bold=True, color="1E293B")
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")
_DATA_FONT = Font(name="Calibri", size=10, color="1F2937")
_ZEBRA_FILL = PatternFill("solid", fgColor="F8FAFC")


def sanitize_sheet_title(title: str, used: set[str]) -> str:
    """Excel limits: <=31 chars, none of []:*?/\\, no edge apostrophes, unique (case-insensitive),
    and the name 'History' is reserved."""
    cleaned = _INVALID_SHEET_CHARS.sub("_", title).strip().strip("'")
    cleaned = cleaned[:MAX_SHEET_TITLE].strip().strip("'") or "Sheet"
    if cleaned.casefold() == "history":
        cleaned = "History_"
    candidate, n = cleaned, 2
    while candidate.casefold() in used:
        suffix = f"_{n}"
        candidate = cleaned[: MAX_SHEET_TITLE - len(suffix)] + suffix
        n += 1
    used.add(candidate.casefold())
    return candidate


class XLSXExporter:
    """Exports extracted tables to Excel.

    All tables are exported onto a unified, beautifully styled primary sheet
    ('All Tables' or 'Table 1') with table title banners, formatted headers,
    and clean spacing. When multiple tables are present, individual sheets
    are also generated for convenient single-table reference.

    Cells are written as strings to prevent formula injection.
    """

    def export(self, document: TableDocument, output_dir: Path) -> list[ExportedFile]:
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.creator = "PDF Table Extraction API"
        used: set[str] = set()

        if len(document.tables) <= 1:
            if document.tables:
                table = document.tables[0]
                sheet = workbook.create_sheet(title=sanitize_sheet_title(f"Table {table.number}", used))
                sheet_widths: dict[int, float] = {}
                self._write_table_block(sheet, table, start_row=1, widths=sheet_widths)
                self._finalize_column_widths(sheet, sheet_widths)
                if table.header_rows:
                    sheet.freeze_panes = sheet.cell(row=table.header_rows + 2, column=1)
            else:
                workbook.create_sheet(title="No Tables")
        else:
            # 1. Primary consolidated sheet containing all tables on one page
            all_sheet = workbook.create_sheet(title=sanitize_sheet_title("All Tables", used))
            current_row = 1
            all_widths: dict[int, float] = {}
            for table in document.tables:
                current_row = self._write_table_block(all_sheet, table, start_row=current_row, widths=all_widths)
                current_row += 2  # 2 blank rows between tables
            self._finalize_column_widths(all_sheet, all_widths)

            # 2. Individual sheets for each table
            for table in document.tables:
                pages = f"p{table.pages[0]}" + (f"-{table.pages[-1]}" if len(table.pages) > 1 else "")
                sheet = workbook.create_sheet(title=sanitize_sheet_title(f"Table {table.number} ({pages})", used))
                sheet_widths = {}
                self._write_table_block(sheet, table, start_row=1, widths=sheet_widths)
                self._finalize_column_widths(sheet, sheet_widths)
                if table.header_rows:
                    sheet.freeze_panes = sheet.cell(row=table.header_rows + 2, column=1)

        workbook.active = 0
        path = output_dir / "extracted-tables.xlsx"
        workbook.save(path)
        return [ExportedFile(path, path.name, XLSX_MEDIA)]

    @classmethod
    def _write_table_block(
        cls,
        sheet: Worksheet,
        table: Table,
        start_row: int,
        widths: dict[int, float],
    ) -> int:
        title_text = table.title or f"Table {table.number}"
        cols = max(table.column_count, 1)

        # 1. Title Banner Row
        sheet.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=start_row,
            end_column=cols,
        )
        title_cell = sheet.cell(row=start_row, column=1, value=title_text)
        title_cell.data_type = "s"
        title_cell.font = _TITLE_FONT
        title_cell.fill = _TITLE_FILL
        title_cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        title_cell.border = _BORDER
        sheet.row_dimensions[start_row].height = 26

        for c in range(2, cols + 1):
            sheet.cell(row=start_row, column=c).border = _BORDER

        # 2. Table Cells (Headers and Data)
        for cell in table.cells:
            r = start_row + 1 + cell.row
            c = cell.column + 1
            text = cell.text[:MAX_CELL_CHARS]
            target = sheet.cell(row=r, column=c, value=text)
            target.data_type = "s"
            target.border = _BORDER

            if cell.is_header:
                target.font = _HEADER_FONT
                target.fill = _HEADER_FILL
                target.alignment = Alignment(
                    wrap_text=True,
                    vertical="center",
                    horizontal="center",
                )
                sheet.row_dimensions[r].height = 22
            else:
                target.font = _DATA_FONT
                target.alignment = Alignment(
                    wrap_text=True,
                    vertical="center",
                    horizontal="center",
                )
                sheet.row_dimensions[r].height = 20
                body_row_idx = cell.row - table.header_rows
                if body_row_idx >= 0 and body_row_idx % 2 == 1:
                    target.fill = _ZEBRA_FILL

            if cell.rowspan > 1 or cell.colspan > 1:
                sheet.merge_cells(
                    start_row=r,
                    start_column=c,
                    end_row=min(r + cell.rowspan - 1, start_row + table.row_count),
                    end_column=min(c + cell.colspan - 1, cols),
                )
            elif text:
                longest = max(len(line) for line in text.split("\n"))
                widths[cell.column] = max(widths.get(cell.column, MIN_COL_WIDTH), longest + 3)

        return start_row + 1 + table.row_count

    @staticmethod
    def _finalize_column_widths(sheet: Worksheet, widths: dict[int, float]) -> None:
        for col_idx, width in widths.items():
            col_letter = get_column_letter(col_idx + 1)
            sheet.column_dimensions[col_letter].width = min(max(width, MIN_COL_WIDTH), MAX_COL_WIDTH)