import io
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pymupdf

from app.config import settings
from app.main import app

client = TestClient(app)


def test_extract_tables_endpoint_returns_json_tables(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)

    document = pymupdf.open()
    page = document.new_page()
    shape = page.new_shape()
    for x in (72, 220, 368):
        shape.draw_line((x, 100), (x, 180))
    for y in (100, 140, 180):
        shape.draw_line((72, y), (368, y))
    shape.finish(color=(0, 0, 0), width=1)
    shape.commit()
    page.insert_text((80, 125), "Item")
    page.insert_text((230, 125), "Price")
    page.insert_text((80, 165), "Tea")
    page.insert_text((230, 165), "5")
    pdf_bytes = document.tobytes()
    document.close()

    response = client.post(
        "/api/v1/pdf/extract-tables",
        params={"output_format": "json", "ocr": "false"},
        files={"file": ("table.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/json")
    assert response.headers["x-table-count"] == "1"
    table = response.json()["tables"][0]
    assert table["columns"] == ["Item", "Price"]
    assert table["rows"] == [{"Item": "Tea", "Price": "5"}]
    assert "Table 1" in table["title"]


def test_extract_borderless_tables_with_title_json(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)

    doc = pymupdf.open()
    p = doc.new_page()
    # Non-border / borderless table with title
    p.insert_text((50, 60), "Table 1: Financial Performance", fontsize=12)
    rows = [
        ("Quarter", "Revenue", "Profit"),
        ("Q1", "10000", "2500"),
        ("Q2", "12000", "3000"),
    ]
    y = 90
    for row in rows:
        x = 50
        for col in row:
            p.insert_text((x, y), col, fontsize=10)
            x += 100
        y += 25

    pdf_bytes = doc.tobytes()
    doc.close()

    response = client.post(
        "/api/v1/pdf/extract-tables",
        params={"output_format": "json", "ocr": "false"},
        files={"file": ("borderless.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    assert response.headers["x-table-count"] == "1"
    data = response.json()
    table = data["tables"][0]
    assert "Financial Performance" in table["title"]
    assert table["columns"] == ["Quarter", "Revenue", "Profit"]
    assert len(table["rows"]) == 2
    assert table["rows"][0]["Quarter"] == "Q1"
    assert table["rows"][1]["Quarter"] == "Q2"


def test_extract_tables_xlsx_all_in_one_page(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)

    doc = pymupdf.open()
    # Page 1
    p1 = doc.new_page()
    p1.insert_text((50, 50), "Table 1: Revenue Breakdown", fontsize=12)
    rows1 = [
        ("Quarter", "Revenue", "Profit"),
        ("Q1", "10000", "2500"),
        ("Q2", "12000", "3000"),
    ]
    y = 75
    for row in rows1:
        x = 50
        for col in row:
            p1.insert_text((x, y), col, fontsize=10)
            x += 100
        y += 22

    # Page 2
    p2 = doc.new_page()
    p2.insert_text((50, 50), "Table 2: Regional Performance", fontsize=12)
    rows2 = [
        ("Region", "Units", "Target"),
        ("North", "500", "450"),
        ("South", "300", "350"),
    ]
    y = 75
    for row in rows2:
        x = 50
        for col in row:
            p2.insert_text((x, y), col, fontsize=10)
            x += 100
        y += 22

    pdf_bytes = doc.tobytes()
    doc.close()

    response = client.post(
        "/api/v1/pdf/extract-tables",
        params={"output_format": "xlsx", "ocr": "false", "merge_tables": "false"},
        files={"file": ("multipage.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["x-table-count"] == "2"

    wb = load_workbook(io.BytesIO(response.content))
    assert "All Tables" in wb.sheetnames
    all_sheet = wb["All Tables"]

    # Verify that both tables are presented on the "All Tables" sheet with title banners
    sheet_text = [
        all_sheet.cell(r, 1).value
        for r in range(1, all_sheet.max_row + 1)
        if all_sheet.cell(r, 1).value
    ]
    assert any("Revenue Breakdown" in val for val in sheet_text)
    assert any("Regional Performance" in val for val in sheet_text)


def test_extract_tables_html_includes_title(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)

    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((50, 60), "Table 1: Customer Accounts", fontsize=12)
    rows = [
        ("Account", "Balance"),
        ("A-100", "500"),
        ("B-200", "750"),
    ]
    y = 90
    for row in rows:
        x = 50
        for col in row:
            p.insert_text((x, y), col, fontsize=10)
            x += 100
        y += 25
    pdf_bytes = doc.tobytes()
    doc.close()

    response = client.post(
        "/api/v1/pdf/extract-tables",
        params={"output_format": "html", "ocr": "false"},
        files={"file": ("accounts.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    assert "Customer Accounts" in response.text
    assert "<h2>" in response.text


def test_extract_tables_markdown_includes_title(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)

    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((50, 60), "Table 1: Inventory Stock", fontsize=12)
    rows = [
        ("Item", "InStock"),
        ("Keyboard", "15"),
        ("Mouse", "30"),
    ]
    y = 90
    for row in rows:
        x = 50
        for col in row:
            p.insert_text((x, y), col, fontsize=10)
            x += 100
        y += 25
    pdf_bytes = doc.tobytes()
    doc.close()

    response = client.post(
        "/api/v1/pdf/extract-tables",
        params={"output_format": "markdown", "ocr": "false"},
        files={"file": ("inventory.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    assert "## Table 1:" in response.text
    assert "Inventory Stock" in response.text


def test_extract_tables_ai_in_table_pdf_no_overdetection(tmp_path, monkeypatch):
    """Test against ai_in_table.pdf: ensures page 1 text is not overdetected as a table,

    and extracts exactly the 3 valid tables (including borderless and multi-page merged).
    """
    monkeypatch.setattr(settings, "OUTPUT_DIR", tmp_path)
    pdf_path = Path("ai_in_table.pdf")
    if not pdf_path.exists():
        pytest.skip("ai_in_table.pdf not found in workspace")

    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    response = client.post(
        "/api/v1/pdf/extract-tables",
        params={"output_format": "json", "merge_tables": "true", "ocr": "false"},
        files={"file": ("ai_in_table.pdf", pdf_bytes, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    data = response.json()
    tables = data["tables"]

    # Must detect exactly 3 tables, no false positive on page 1
    assert len(tables) == 3, f"Expected 3 tables, got {len(tables)}"

    # Check pages
    assert tables[0]["pages"] == [5]
    assert tables[1]["pages"] == [6]
    assert tables[2]["pages"] == [6, 7]

    # Check headers and clean text (no zero-width spaces \u200b)
    t1_headers = [c["text"] for c in tables[0]["cells"] if c["row"] == 0]
    assert t1_headers == ["Name", "Roll", "Age", "Address", "Choices"]

    t2_headers = [c["text"] for c in tables[1]["cells"] if c["row"] == 0]
    assert t2_headers == ["City", "Code", "People"]

    t3_headers = [c["text"] for c in tables[2]["cells"] if c["row"] == 0]
    assert t3_headers == ["City", "Code", "People"]
    assert tables[2]["row_count"] == 11

    # Ensure zero-width spaces are stripped across all cells
    for tbl in tables:
        for c in tbl["cells"]:
            assert "\u200b" not in c["text"]
            assert "\ufeff" not in c["text"]

